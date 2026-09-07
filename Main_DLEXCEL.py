"""
Sensor Logger - PC-side WiFi receiver
---------------------------------------
The Pico sends a reading once a minute, always. This server:
  - Always shows the latest reading on the dashboard (live view)
  - Only WRITES readings to Excel while a sleep session is active
    (started by "Going to Sleep", ended by "I'm Awake")
  - Computes an average roughly every hour during a session, plus a
    final partial average when the session ends
  - Logs each session's start/end time and your sleep-quality answer

First-time setup:
    pip install flask openpyxl

Run:
    python server.py
"""

import datetime

from flask import Flask, request, jsonify, render_template_string, send_file
from openpyxl import Workbook, load_workbook

# ------------- CONFIGURE THIS -------------
EXCEL_PATH = r"C:\Users\Khoa\Desktop\sensor_log.xlsx"
# -------------------------------------------

RAW_SHEET = "Sensor Data"
HOURLY_SHEET = "Hourly Averages"
SLEEP_SHEET = "Sleep Log"

RAW_HEADER = ["timestamp", "air_quality", "lux", "temperature_c", "humidity_pct", "sound_level"]
HOURLY_HEADER = ["period_start", "period_end", "avg_air_quality", "avg_lux",
                  "avg_temperature_c", "avg_humidity_pct", "avg_sound_level", "sample_count"]
SLEEP_HEADER = ["session_start", "session_end", "slept_well"]

app = Flask(__name__)

# ---- In-memory session state ----
logging_active = False
session_start = None
last_hourly_calc = None
latest_reading = None  # always updated, shown on dashboard regardless of session state


def get_or_create_workbook(path):
    try:
        wb = load_workbook(path)
    except FileNotFoundError:
        wb = Workbook()
        wb.remove(wb.active)

    if RAW_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(RAW_SHEET)
        ws.append(RAW_HEADER)
    if HOURLY_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(HOURLY_SHEET)
        ws.append(HOURLY_HEADER)
    if SLEEP_SHEET not in wb.sheetnames:
        ws = wb.create_sheet(SLEEP_SHEET)
        ws.append(SLEEP_HEADER)

    wb.save(path)
    return wb


def compute_average(wb, period_start, period_end):
    """Average raw rows within [period_start, period_end) and append to
    the Hourly Averages sheet. Returns True if a row was added."""
    ws_raw = wb[RAW_SHEET]

    sums = {"air_quality": 0.0, "lux": 0.0, "temperature_c": 0.0, "humidity_pct": 0.0, "sound_level": 0.0}
    counts = {k: 0 for k in sums}

    for row in ws_raw.iter_rows(min_row=2, values_only=True):
        ts_raw = row[0]
        ts = datetime.datetime.fromisoformat(ts_raw) if isinstance(ts_raw, str) else ts_raw
        if ts is None or not (period_start <= ts < period_end):
            continue
        for i, key in enumerate(["air_quality", "lux", "temperature_c", "humidity_pct", "sound_level"], start=1):
            val = row[i]
            if val is not None and val != "":
                sums[key] += float(val)
                counts[key] += 1

    sample_count = max(counts.values()) if counts else 0
    if sample_count == 0:
        return False

    def avg(key):
        return round(sums[key] / counts[key], 2) if counts[key] > 0 else None

    ws_hourly = wb[HOURLY_SHEET]
    ws_hourly.append([
        period_start.isoformat(timespec="seconds"),
        period_end.isoformat(timespec="seconds"),
        avg("air_quality"),
        avg("lux"),
        avg("temperature_c"),
        avg("humidity_pct"),
        avg("sound_level"),
        sample_count,
    ])
    return True


def maybe_compute_hourly_average(wb):
    global last_hourly_calc
    now = datetime.datetime.now()
    if last_hourly_calc is None or (now - last_hourly_calc).total_seconds() < 3600:
        return
    compute_average(wb, last_hourly_calc, now)
    last_hourly_calc = now


DASHBOARD_HTML = """
<!DOCTYPE html>
<html>
<head>
  <title>Sensor Log Dashboard</title>
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <style>
    body { font-family: Arial, sans-serif; background: #111; color: #eee; margin: 0; padding: 24px; }
    h1 { font-size: 20px; margin-bottom: 4px; }
    h2 { font-size: 16px; margin-top: 32px; color: #ccc; }
    #status { color: #7cf; font-size: 13px; margin-bottom: 16px; }
    table { border-collapse: collapse; width: 100%; max-width: 900px; }
    th, td { padding: 8px 12px; text-align: left; border-bottom: 1px solid #333; font-size: 14px; }
    th { background: #1e1e1e; position: sticky; top: 0; }
    tr:hover { background: #1a1a1a; }
    .btn {
      margin: 8px 8px 8px 0; padding: 12px 24px; font-size: 15px; font-weight: bold;
      border: none; border-radius: 6px; cursor: pointer; color: #fff;
    }
    #sleepBtn { background: #57c; }
    #sleepBtn:hover { background: #68d; }
    #wakeBtn { background: #3a7; }
    #wakeBtn:hover { background: #4b8; }
    .btn:disabled { background: #444; cursor: not-allowed; }
    #sessionStatus { margin: 8px 0 20px; font-size: 13px; color: #7cf; }
    #liveBox {
      background: #1a1a1a; border-radius: 8px; padding: 16px; max-width: 500px;
      margin-bottom: 20px; font-size: 14px; line-height: 1.6;
    }
  </style>
</head>
<body>
  <h1>Sensor Log</h1>
  <div id="status">Loading...</div>
  <a href="/download" class="btn" style="background:#888; text-decoration:none; display:inline-block;">Download as Excel</a>

  <div id="liveBox">Waiting for first reading...</div>

  <button id="sleepBtn" class="btn" onclick="goToSleep()">Going to Sleep</button>
  <button id="wakeBtn" class="btn" onclick="wakeUp()">I'm Awake</button>
  <div id="sessionStatus"></div>

  <h2>Logged Readings (this session / recent sessions)</h2>
  <table>
    <thead>
      <tr><th>Timestamp</th><th>Air Quality</th><th>Lux</th><th>Temp (C)</th><th>Humidity (%)</th><th>Sound</th></tr>
    </thead>
    <tbody id="rows"></tbody>
  </table>

  <h2>Hourly Averages</h2>
  <table>
    <thead>
      <tr><th>Period</th><th>Avg Air</th><th>Avg Lux</th><th>Avg Temp</th><th>Avg Humidity</th><th>Avg Sound</th><th>Samples</th></tr>
    </thead>
    <tbody id="hourlyRows"></tbody>
  </table>

  <script>
    async function goToSleep() {
      const res = await fetch('/sleep/start', { method: 'POST' });
      const data = await res.json();
      document.getElementById('sessionStatus').textContent =
        'Sleep session started at ' + new Date().toLocaleString() + '. Logging now.';
      updateButtons(true);
    }

    async function wakeUp() {
      const sleptWell = confirm("Did you sleep well?\\n\\nOK = Yes\\nCancel = No");
      const res = await fetch('/sleep/end', {
        method: 'POST',
        headers: {'Content-Type': 'application/json'},
        body: JSON.stringify({slept_well: sleptWell})
      });
      if (res.ok) {
        document.getElementById('sessionStatus').textContent =
          'Session ended. Slept well = ' + (sleptWell ? 'Yes' : 'No') + '. Logging stopped.';
        updateButtons(false);
      } else {
        document.getElementById('sessionStatus').textContent = 'Failed to record (was a session even active?).';
      }
    }

    function updateButtons(active) {
      document.getElementById('sleepBtn').disabled = active;
      document.getElementById('wakeBtn').disabled = !active;
    }

    async function refresh() {
      try {
        const res = await fetch('/api/data');
        const data = await res.json();

        // Live reading box - always shown, regardless of logging state
        const live = data.live;
        const liveBox = document.getElementById('liveBox');
        if (live) {
          liveBox.innerHTML =
            `<b>Live reading</b> (${live.timestamp})<br>` +
            `Air Quality: ${live.air_quality ?? '-'} &nbsp; ` +
            `Lux: ${live.lux ?? '-'} &nbsp; ` +
            `Temp: ${live.temperature_c ?? '-'} C &nbsp; ` +
            `Humidity: ${live.humidity_pct ?? '-'}% &nbsp; ` +
            `Sound: ${live.sound_level ?? '-'}` +
            `<br><span style="color:${data.logging_active ? '#8f8' : '#f88'}">` +
            (data.logging_active ? 'Logging: ACTIVE (sleep session in progress)' : 'Logging: OFF (not in a sleep session)') +
            `</span>`;
        }
        updateButtons(data.logging_active);

        const tbody = document.getElementById('rows');
        tbody.innerHTML = '';
        data.rows.forEach(row => {
          const tr = document.createElement('tr');
          tr.innerHTML = row.map(cell => `<td>${cell === null || cell === '' ? '-' : cell}</td>`).join('');
          tbody.appendChild(tr);
        });

        const hourlyBody = document.getElementById('hourlyRows');
        hourlyBody.innerHTML = '';
        data.hourly.forEach(row => {
          const tr = document.createElement('tr');
          tr.innerHTML = row.map(cell => `<td>${cell === null || cell === '' ? '-' : cell}</td>`).join('');
          hourlyBody.appendChild(tr);
        });

        document.getElementById('status').textContent =
          `Showing latest ${data.rows.length} of ${data.total} logged readings — updated ${new Date().toLocaleString()}`;
      } catch (e) {
        document.getElementById('status').textContent = 'Failed to load data: ' + e;
      }
    }
    refresh();
    setInterval(refresh, 5000);
  </script>
</body>
</html>
"""


@app.route("/")
def dashboard():
    return render_template_string(DASHBOARD_HTML)


@app.route("/download")
def download_excel():
    get_or_create_workbook(EXCEL_PATH)  # make sure it exists before trying to send it
    filename = f"sensor_log_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        EXCEL_PATH,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.route("/api/data")
def api_data():
    wb = get_or_create_workbook(EXCEL_PATH)

    ws_raw = wb[RAW_SHEET]
    all_rows = list(ws_raw.iter_rows(min_row=2, values_only=True))
    latest = list(reversed(all_rows[-50:]))

    ws_hourly = wb[HOURLY_SHEET]
    hourly_rows = list(ws_hourly.iter_rows(min_row=2, values_only=True))
    hourly_latest = list(reversed(hourly_rows[-24:]))

    return jsonify({
        "rows": latest,
        "total": len(all_rows),
        "hourly": hourly_latest,
        "live": latest_reading,
        "logging_active": logging_active,
    })


@app.route("/log", methods=["POST"])
def log_reading():
    global latest_reading

    data = request.get_json()
    if not data:
        return jsonify({"error": "no JSON body received"}), 400

    timestamp = datetime.datetime.now().isoformat(timespec="seconds")

    # Always update the live reading, whether or not we're logging
    latest_reading = {
        "timestamp": timestamp,
        "air_quality": data.get("air_quality"),
        "lux": data.get("lux"),
        "temperature_c": data.get("temperature_c"),
        "humidity_pct": data.get("humidity_pct"),
        "sound_level": data.get("sound_level"),
    }

    if not logging_active:
        print(f"Received (not logging): {latest_reading}")
        return jsonify({"status": "ok", "logged": False}), 200

    row = [
        timestamp,
        data.get("air_quality"),
        data.get("lux"),
        data.get("temperature_c"),
        data.get("humidity_pct"),
        data.get("sound_level"),
    ]

    wb = get_or_create_workbook(EXCEL_PATH)
    ws = wb[RAW_SHEET]
    ws.append(row)

    maybe_compute_hourly_average(wb)

    wb.save(EXCEL_PATH)

    print(f"Logged: {row}")
    return jsonify({"status": "ok", "logged": True}), 200


@app.route("/sleep/start", methods=["POST"])
def sleep_start():
    global logging_active, session_start, last_hourly_calc
    logging_active = True
    session_start = datetime.datetime.now()
    last_hourly_calc = session_start
    print(f"Sleep session started at {session_start.isoformat(timespec='seconds')}")
    return jsonify({"status": "ok", "session_start": session_start.isoformat(timespec="seconds")}), 200


@app.route("/sleep/end", methods=["POST"])
def sleep_end():
    global logging_active, session_start, last_hourly_calc

    if not logging_active or session_start is None:
        return jsonify({"error": "no active session"}), 400

    data = request.get_json()
    if data is None or "slept_well" not in data:
        return jsonify({"error": "missing slept_well"}), 400

    session_end = datetime.datetime.now()
    answer = "Yes" if data["slept_well"] else "No"

    wb = get_or_create_workbook(EXCEL_PATH)

    # Final partial-hour average for whatever time is left since the last average
    compute_average(wb, last_hourly_calc, session_end)

    ws_sleep = wb[SLEEP_SHEET]
    ws_sleep.append([
        session_start.isoformat(timespec="seconds"),
        session_end.isoformat(timespec="seconds"),
        answer,
    ])
    wb.save(EXCEL_PATH)

    print(f"Sleep session ended: {session_start} -> {session_end}, slept_well={answer}")

    logging_active = False
    session_start = None
    last_hourly_calc = None

    return jsonify({"status": "ok"}), 200


if __name__ == "__main__":
    print(f"Logging to: {EXCEL_PATH}")
    get_or_create_workbook(EXCEL_PATH)
    print("Starting server on http://0.0.0.0:5000 ...")
    print("Waiting for the Pico to send data. Press Ctrl+C to stop.\n")
    app.run(host="0.0.0.0", port=5000)