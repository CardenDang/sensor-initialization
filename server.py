"""
Sensor Logger - PC-side WiFi receiver
---------------------------------------
Runs a small web server on your PC. The Pico sends an HTTP POST request
to this server every second with its sensor readings, and this script
appends each one as a row in an Excel (.xlsx) file.

First-time setup (run once in a terminal):
    pip install flask openpyxl

Run:
    python server.py

Then find your PC's local IP address (run "ipconfig" in Command Prompt,
look for "IPv4 Address" e.g. 192.168.1.50) and put that into the
SERVER_URL in the Pico's main.py, e.g.:
    SERVER_URL = "http://MY_IP_HERE:5000/log"

Your PC and the Pico must be on the SAME WiFi network for this to work.
"""

import datetime

from flask import Flask, request, jsonify, render_template_string
from openpyxl import Workbook, load_workbook

# ------------- CONFIGURE THIS -------------
# <-- change to wherever you want the file
EXCEL_PATH = r"C:\Users\Khoa\Desktop\sensor_log.xlsx"
# -------------------------------------------

HEADER = ["timestamp", "air_quality", "lux",
          "temperature_c", "humidity_pct", "sound_level"]

app = Flask(__name__)


def get_or_create_workbook(path):
    try:
        wb = load_workbook(path)
        ws = wb.active
        return wb, ws
    except FileNotFoundError:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sensor Data"
        ws.append(HEADER)
        wb.save(path)
        return wb, ws


DASHBOARD_HTML = """
<!DOCTYPE html>
<html>
<head>
  <title>Sensor Log Dashboard</title>
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <style>
    body { font-family: Arial, sans-serif; background: #111; color: #eee; margin: 0; padding: 24px; }
    h1 { font-size: 20px; margin-bottom: 4px; }
    #status { color: #7cf; font-size: 13px; margin-bottom: 16px; }
    table { border-collapse: collapse; width: 100%; max-width: 900px; }
    th, td { padding: 8px 12px; text-align: left; border-bottom: 1px solid #333; font-size: 14px; }
    th { background: #1e1e1e; position: sticky; top: 0; }
    tr:hover { background: #1a1a1a; }
    tr:first-child { color: #8f8; font-weight: bold; }
  </style>
</head>
<body>
  <h1>Sensor Log</h1>
  <div id="status">Loading...</div>
  <table>
    <thead>
      <tr>
        <th>Timestamp</th><th>Air Quality</th><th>Lux</th><th>Temp (C)</th><th>Humidity (%)</th><th>Sound</th>
      </tr>
    </thead>
    <tbody id="rows"></tbody>
  </table>

  <script>
    async function refresh() {
      try {
        const res = await fetch('/api/data');
        const data = await res.json();
        const tbody = document.getElementById('rows');
        tbody.innerHTML = '';
        data.rows.forEach(row => {
          const tr = document.createElement('tr');
          tr.innerHTML = row.map(cell => `<td>${cell === null || cell === '' ? '-' : cell}</td>`).join('');
          tbody.appendChild(tr);
        });
        document.getElementById('status').textContent =
          `Showing latest ${data.rows.length} of ${data.total} readings — updated ${new Date().toLocaleTimeString()}`;
      } catch (e) {
        document.getElementById('status').textContent = 'Failed to load data: ' + e;
      }
    }
    refresh();
    setInterval(refresh, 2000); // auto-refresh every 2 seconds
  </script>
</body>
</html>
"""


@app.route("/")
def dashboard():
    return render_template_string(DASHBOARD_HTML)


@app.route("/api/data")
def api_data():
    try:
        wb = load_workbook(EXCEL_PATH)
        ws = wb.active
        all_rows = list(ws.iter_rows(
            min_row=2, values_only=True))  # skip header
        latest = all_rows[-50:]      # only send the most recent 50 rows
        latest.reverse()             # newest first
        return jsonify({"rows": latest, "total": len(all_rows)})
    except FileNotFoundError:
        return jsonify({"rows": [], "total": 0})


@app.route("/log", methods=["POST"])
def log_reading():
    data = request.get_json()
    if not data:
        return jsonify({"error": "no JSON body received"}), 400

    timestamp = datetime.datetime.now().isoformat(timespec="seconds")
    row = [
        timestamp,
        data.get("air_quality"),
        data.get("lux"),
        data.get("temperature_c"),
        data.get("humidity_pct"),
        data.get("sound_level"),
    ]

    wb, ws = get_or_create_workbook(EXCEL_PATH)
    ws.append(row)
    wb.save(EXCEL_PATH)

    print(f"Logged: {row}")
    return jsonify({"status": "ok"}), 200


if __name__ == "__main__":
    print(f"Logging to: {EXCEL_PATH}")
    print("Starting server on http://0.0.0.0:5000 ...")
    print("Waiting for the Pico to send data. Press Ctrl+C to stop.\n")
    app.run(host="0.0.0.0", port=5000)
